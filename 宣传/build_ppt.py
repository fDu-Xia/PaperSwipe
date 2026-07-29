# -*- coding: utf-8 -*-
"""
生成 PaperSwipe 项目介绍 PPT
用法: python3 build_ppt.py
输出: 宣传/PaperSwipe-项目介绍.pptx
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image

BASE = "/Users/buluu/Documents/PaperSwipe"
ASSETS = os.path.join(BASE, "webassets")
INTROWEB_IMG = os.path.join(BASE, "introweb", "images")
OUT = os.path.join(BASE, "宣传", "PaperSwipe-项目介绍.pptx")

FONT = "PingFang SC"

# ---------- 品牌色板（与 introweb 官网一致） ----------
BG = RGBColor(0xF4, 0xF5, 0xF7)
BG_ALT = RGBColor(0xEC, 0xEE, 0xF2)
INK = RGBColor(0x11, 0x13, 0x1A)
INK_DIM = RGBColor(0x56, 0x5A, 0x68)
INK_FAINT = RGBColor(0x8A, 0x8E, 0x9C)
VIOLET = RGBColor(0x7C, 0x6C, 0xFF)
VIOLET_STRONG = RGBColor(0x66, 0x57, 0xE8)
CYAN = RGBColor(0x2F, 0xD0, 0xC9)
CORAL = RGBColor(0xFF, 0x7A, 0x63)
YELLOW = RGBColor(0xF6, 0xC4, 0x53)
GREEN = RGBColor(0x3E, 0xCF, 0x8E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xDF, 0xE2, 0xE8)
DARK_PANEL = RGBColor(0x17, 0x18, 0x22)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]

PAGE = {"n": 0}


def new_slide(bg=BG):
    slide = prs.slides.add_slide(BLANK)
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    rect.fill.solid()
    rect.fill.fore_color.rgb = bg
    rect.line.fill.background()
    rect.shadow.inherit = False
    PAGE["n"] += 1
    if PAGE["n"] > 1:
        add_footer(slide)
    return slide


def add_footer(slide):
    box = slide.shapes.add_textbox(Inches(0.6), SH - Inches(0.5), Inches(6), Inches(0.35))
    tf = box.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "PaperSwipe · 别再搜论文了，开始刷论文"
    set_font(r, size=10, color=INK_FAINT)
    box2 = slide.shapes.add_textbox(SW - Inches(1.4), SH - Inches(0.5), Inches(0.9), Inches(0.35))
    tf2 = box2.text_frame
    p2 = tf2.paragraphs[0]
    p2.alignment = PP_ALIGN.RIGHT
    r2 = p2.add_run()
    r2.text = str(PAGE["n"])
    set_font(r2, size=10, color=INK_FAINT)


def set_font(run, size=16, color=INK, bold=False, italic=False, name=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = name
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    ea = rPr.makeelement(qn('a:ea'), {'typeface': name})
    rPr.append(ea)


def add_rect(slide, left, top, width, height, color, line=False, line_color=None, radius=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, left, top, width, height)
    if radius:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    if color is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = color
    if line and line_color:
        shp.line.color.rgb = line_color
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def add_text(slide, left, top, width, height, runs, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, line_spacing=1.2, space_after=8, wrap=True):
    """runs: list[list[tuple(text, size, color, bold)]] -> 每个子list是一段"""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        for (t, size, color, bold) in para:
            r = p.add_run()
            r.text = t
            set_font(r, size=size, color=color, bold=bold)
    return box


def add_kicker(slide, text, top=Inches(0.55), color=VIOLET, left=Inches(0.75)):
    bar = add_rect(slide, left, top + Inches(0.09), Inches(0.28), Pt(3), color)
    add_text(slide, left + Inches(0.38), top, Inches(6), Inches(0.4),
              [[(text, 13, color, True)]])


def add_title(slide, text, top=Inches(0.95), size=30, color=INK, left=Inches(0.75), width=Inches(11.5)):
    add_text(slide, left, top, width, Inches(1.0), [[(text, size, color, True)]], line_spacing=1.15)


def bullets(slide, items, left, top, width, height, size=14.5, color=INK_DIM, gap=10, mark="•", mark_color=None):
    runs = []
    for it in items:
        runs.append([(f"{mark}  ", size, mark_color or VIOLET, True), (it, size, color, False)])
    add_text(slide, left, top, width, height, runs, space_after=gap, line_spacing=1.25)


def picture_fit(slide, path, left, top, max_w, max_h, border=True, border_color=None):
    im = Image.open(path)
    iw, ih = im.size
    ratio = min(max_w / iw, max_h / ih)
    w, h = int(iw * ratio), int(ih * ratio)
    pic_left = left + int((max_w - w) / 2)
    pic_top = top + int((max_h - h) / 2)
    if border:
        pad = Emu(int(0.08 * 914400))
        add_rect(slide, pic_left - pad, pic_top - pad, w + pad * 2, h + pad * 2,
                 WHITE, radius=0.06)
    pic = slide.shapes.add_picture(path, pic_left, pic_top, width=w, height=h)
    pic.line.color.rgb = border_color or LINE
    pic.line.width = Pt(0.75) if border else Pt(0)
    return pic


def table_slide(slide, headers, rows, left, top, width, height, col_widths=None,
                 header_bg=INK, header_color=WHITE, body_size=12.5, header_size=13):
    n_rows = len(rows) + 1
    n_cols = len(headers)
    gframe = slide.shapes.add_table(n_rows, n_cols, left, top, width, height)
    table = gframe.table
    if col_widths:
        total = sum(col_widths)
        for i, cw in enumerate(col_widths):
            table.columns[i].width = int(width * cw / total)
    for j, htext in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = ""
        tf = cell.text_frame
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = htext
        set_font(r, size=header_size, color=header_color, bold=True)
        cell.fill.solid()
        cell.fill.fore_color.rgb = header_bg
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_top = Pt(6); cell.margin_bottom = Pt(6)
        cell.margin_left = Pt(8); cell.margin_right = Pt(8)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.cell(i + 1, j)
            cell.text = ""
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = val
            bold0 = (j == 0)
            set_font(r, size=body_size, color=INK if j == 0 else INK_DIM, bold=bold0)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if i % 2 == 0 else BG_ALT
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_top = Pt(6); cell.margin_bottom = Pt(6)
            cell.margin_left = Pt(8); cell.margin_right = Pt(8)
    return table


# =====================================================================
# 1. 封面
# =====================================================================
s = new_slide()
add_rect(s, 0, 0, Inches(0.18), SH, VIOLET)
logo_path = os.path.join(BASE, "logo.png")
picture_fit(s, logo_path, Inches(0.9), Inches(1.0), Inches(1.1), Inches(1.1), border=False)
add_text(s, Inches(2.2), Inches(1.05), Inches(9), Inches(1.0),
         [[("PaperSwipe", 30, INK, True)]])
add_text(s, Inches(2.2), Inches(1.62), Inches(9), Inches(0.5),
         [[("AI 驱动的论文发现与阅读管理工具", 14, INK_FAINT, False)]])

add_text(s, Inches(0.9), Inches(3.15), Inches(11.5), Inches(2.0),
         [[("别再", 44, INK, True), ("搜", 44, CORAL, True), ("论文了，", 44, INK, True)],
          [("开始", 44, INK, True), ("刷", 44, VIOLET, True), ("论文。", 44, INK, True)]],
         line_spacing=1.08)

add_text(s, Inches(0.9), Inches(5.15), Inches(10.5), Inches(0.9),
         [[("What if Google Scholar felt like TikTok?", 15, INK_DIM, False)]])

add_text(s, Inches(0.9), SH - Inches(0.85), Inches(8), Inches(0.4),
         [[("项目介绍 · PaperSwipe Team", 12, INK_FAINT, False)]])

# =====================================================================
# 2. 痛点：读论文最贵的成本
# =====================================================================
s = new_slide()
add_kicker(s, "痛点 · PAIN POINT")
add_title(s, "读论文最贵的成本，从来不是「读」，而是「找」")

steps = [
    "打开 Google Scholar，搜一个关键词，出来几千篇结果",
    "打开第一篇，翻译摘要，看完关掉；打开第二篇……如此重复几十次",
    "终于收藏了 20 篇论文，标签打好、文件夹建好，感觉自己很努力",
    "最终真正读完的：0 篇",
]
top = Inches(1.9)
for i, txt in enumerate(steps):
    y = top + Inches(0.62) * i
    is_last = i == len(steps) - 1
    circ = add_rect(s, Inches(0.85), y, Inches(0.4), Inches(0.4),
                     CORAL if is_last else WHITE, line=not is_last, line_color=LINE, radius=1.0)
    add_text(s, Inches(0.85), y, Inches(0.4), Inches(0.4),
             [[(f"{i+1}", 13, WHITE if is_last else INK_FAINT, True)]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(1.45), y + Inches(0.02), Inches(10.6), Inches(0.5),
             [[(txt, 15, INK if is_last else INK_DIM, is_last)]], anchor=MSO_ANCHOR.MIDDLE)

add_text(s, Inches(0.85), Inches(4.7), Inches(11.3), Inches(0.9),
         [[("问题从来不是论文不够多，而是从「几千篇结果」到「值得读的那一篇」之间，", 14, INK_DIM, False)],
          [("隔着太多重复、低效、几乎没有信息增量的体力劳动。", 14, INK_DIM, False)]])

# =====================================================================
# 3. 痛点拆解（五个子问题）
# =====================================================================
s = new_slide()
add_kicker(s, "痛点拆解")
add_title(s, "被现有工具系统性忽视的五个问题")

pains = [
    ("信息过载，检索低效", "关键词式搜索召回几千篇结果，仍需逐篇点开详情页、翻看摘要才能判断相关性，交互方式老套繁琐，认知负荷极高", CORAL),
    ("摘要阅读门槛高、耗时长", "英文摘要专业术语密集，且主要面向论文 reviewer 撰写，未必直接体现论文的真实核心贡献；非母语用户还需反复翻译、追问细节，5~10 分钟才能判断一篇论文是否值得深读", YELLOW),
    ("收藏即遗忘", "文献管理软件擅长「管理已收藏的」，但对「选出来之后什么时候读、按什么顺序读」完全没有解决方案，收藏夹变成论文坟场", VIOLET),
    ("热点感知滞后", "需要每天刷 X/Twitter、逛 arXiv 才能感知领域动态，个人很难持续跟踪全球研究前沿，容易错过关键工作或重复造轮子", CYAN),
    ("碎片化时间难以利用", "精读离不开 PC 端整块时间，但「筛选值不值得读」完全可以用通勤、排队等碎片时间在手机上完成——现有工具却全是桌面形态，筛选反而挤占了宝贵的整块阅读时间", GREEN),
]
top = Inches(1.85)
col_h = Inches(0.98)
for i, (title, desc, color) in enumerate(pains):
    y = top + col_h * i
    add_rect(s, Inches(0.75), y + Inches(0.06), Inches(0.09), Inches(0.72), color)
    add_text(s, Inches(1.05), y, Inches(2.6), col_h,
             [[(title, 15, INK, True)]], anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(3.75), y, Inches(8.7), col_h,
             [[(desc, 12.5, INK_DIM, False)]], anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.2)

# =====================================================================
# 4. 大胆的想法
# =====================================================================
s = new_slide(bg=DARK_PANEL)
add_text(s, Inches(1.2), Inches(2.1), Inches(11), Inches(0.5),
         [[("一个大胆的想法", 14, CYAN, True)]])
add_text(s, Inches(1.2), Inches(2.7), Inches(11), Inches(1.8),
         [[("如果找论文，", 36, WHITE, True)],
          [("能像刷 ", 36, WHITE, True), ("TikTok", 36, VIOLET, True), (" 一样呢？", 36, WHITE, True)]],
         line_spacing=1.15)
add_text(s, Inches(1.2), Inches(4.35), Inches(10.5), Inches(1.6),
         [[("不用你输入关键词、翻几十页搜索结果、逐篇打开判断相关性。", 15, RGBColor(0xC7,0xC9,0xD3), False)],
          [("AI 先理解你是谁、最近在研究什么、已经读过哪些论文，", 15, RGBColor(0xC7,0xC9,0xD3), False)],
          [("然后直接把你最可能感兴趣的论文推到你面前——你只需要滑一下。", 15, RGBColor(0xC7,0xC9,0xD3), False)]],
         line_spacing=1.4)
add_text(s, Inches(1.2), Inches(5.85), Inches(8), Inches(0.5),
         [[("于是有了 ", 18, RGBColor(0xC7,0xC9,0xD3), False), ("PaperSwipe", 18, WHITE, True), (" 🎉", 18, WHITE, False)]])

# =====================================================================
# 5. 产品方案总览
# =====================================================================
s = new_slide()
add_kicker(s, "产品方案")
add_title(s, "四个模块，闭环覆盖「发现 → 判断 → 归档 → 阅读 → 追踪热点」")

modules = [
    ("01", "Explore", "AI 认识你，主动发现论文", VIOLET),
    ("02", "Library", "划走的论文，有了真正的家", CYAN),
    ("03", "Schedule", "AI 帮你规划阅读，配合对话助手", CORAL),
    ("04", "Trending", "今天，全世界都在看什么", GREEN),
]
card_w = Inches(2.75)
gap = Inches(0.25)
left0 = Inches(0.75)
top = Inches(2.3)
card_h = Inches(3.6)
for i, (num, name, desc, color) in enumerate(modules):
    x = left0 + (card_w + gap) * i
    add_rect(s, x, top, card_w, card_h, WHITE, radius=0.06, line=True, line_color=LINE)
    add_rect(s, x, top, card_w, Inches(0.12), color)
    add_text(s, x + Inches(0.25), top + Inches(0.35), card_w - Inches(0.5), Inches(0.5),
             [[(num, 22, color, True)]])
    add_text(s, x + Inches(0.25), top + Inches(0.95), card_w - Inches(0.5), Inches(0.6),
             [[(name, 20, INK, True)]])
    add_text(s, x + Inches(0.25), top + Inches(1.55), card_w - Inches(0.5), Inches(1.6),
             [[(desc, 13, INK_DIM, False)]], line_spacing=1.3)

# =====================================================================
# 6. Explore
# =====================================================================
s = new_slide()
add_kicker(s, "01 · Explore")
add_title(s, "AI 主动发现论文，四向滑动 5 秒完成判断")
bullets(s, [
    "AI 自动把研究方向拆解为精准检索词，实时联通 Semantic Scholar / arXiv / OpenAlex 等全球权威学术源进行检索",
    "每篇论文生成双面知识卡片：正面一句话讲清核心贡献，背面拆解问题、方法、结果、创新点与推荐理由",
    "四向滑动完成决策（左滑跳过、右滑收藏、上滑标重点、下滑标已读），同时支持按钮和方向键操作",
    "结合用户画像与历史滑动行为，通过推荐算法持续迭代优化结果，越用越懂你的研究方向",
], Inches(0.75), Inches(1.95), Inches(6.5), Inches(4.2), size=14, gap=14)

picture_fit(s, os.path.join(ASSETS, "explore/pics/explore1-卡片正面.png"),
            Inches(7.55), Inches(1.55), Inches(2.5), Inches(5.1))
picture_fit(s, os.path.join(ASSETS, "explore/pics/explore2-卡片背面.png"),
            Inches(10.15), Inches(1.55), Inches(2.5), Inches(5.1))

# =====================================================================
# 7. Library
# =====================================================================
s = new_slide()
add_kicker(s, "02 · Library")
add_title(s, "划走的每一篇论文，有了真正属于你的家")
bullets(s, [
    "三栏自动分类（感兴趣 / 重点 / 已读），阅读进度一目了然",
    "AI 自动提炼主题标签，支持标签筛选，秒速定位历史论文",
    "长按卡片可快速归档、删除，或一键安排进阅读日程",
    "支持 BibTeX 导出与 Zotero 联动，无缝接入已有文献工作流",
], Inches(0.75), Inches(1.95), Inches(6.5), Inches(4.2), size=14, gap=14)

picture_fit(s, os.path.join(ASSETS, "library/pics/library首页.png"),
            Inches(7.7), Inches(1.35), Inches(2.6), Inches(5.4))
picture_fit(s, os.path.join(ASSETS, "library/pics/library筛选topic.png"),
            Inches(10.35), Inches(1.35), Inches(2.6), Inches(5.4))

# =====================================================================
# 8. Schedule + AI 助手
# =====================================================================
s = new_slide()
add_kicker(s, "03 · Schedule & AI 助手")
add_title(s, "AI 帮你规划阅读，论文不再躺在收藏夹吃灰")
bullets(s, [
    "AI 自动规划阅读周计划，长按拖拽即可跨日期调整",
    "AI 对话助手与 Schedule 联通，可根据研究目标、DDL 和可投入时间自动排期",
    "直接提问「今天最值得读哪篇？」，AI 秒回答案和理由",
    "提问「帮我规划这个方向的阅读 roadmap」，AI 排出成体系的学习路径",
], Inches(0.75), Inches(1.95), Inches(6.5), Inches(4.2), size=14, gap=14)

picture_fit(s, os.path.join(ASSETS, "schedule/pics/schedule首页.png"),
            Inches(7.7), Inches(1.35), Inches(2.6), Inches(5.4))
picture_fit(s, os.path.join(ASSETS, "robot/pics/robot对话例3-为我规划读论文roadmap.png"),
            Inches(10.35), Inches(1.35), Inches(2.6), Inches(5.4))

# =====================================================================
# 9. Trending
# =====================================================================
s = new_slide()
add_kicker(s, "04 · Trending")
add_title(s, "今天，全世界都在看什么")
bullets(s, [
    "实时全球研究热点榜，支持切换查看「本领域」或「全部领域」当下最火的 Topic",
    "一键展开该 Topic 下最具代表性、最火的论文",
    "重要论文引用飙升时第一时间提醒，永远站在信息最前沿",
], Inches(0.75), Inches(1.95), Inches(6.5), Inches(4.2), size=14, gap=14)

picture_fit(s, os.path.join(ASSETS, "trending/pics/trending首页.png"),
            Inches(7.7), Inches(1.35), Inches(2.6), Inches(5.4))
picture_fit(s, os.path.join(ASSETS, "trending/pics/trending-展开最火topic的五篇论文.png"),
            Inches(10.35), Inches(1.35), Inches(2.6), Inches(5.4))

# =====================================================================
# 10. 技术架构与产品形态
# =====================================================================
s = new_slide()
add_kicker(s, "技术架构")
add_title(s, "工程效率优先的技术选型")

add_text(s, Inches(0.75), Inches(1.9), Inches(11.5), Inches(0.8),
         [[("当前形态：", 14, INK, True),
           ("手机优先、零构建的 Web H5（打开即用，无需安装）；已基于 Capacitor 完成 iOS 端封装，", 14, INK_DIM, False)],
          [("可在 iPhone 真机上直接调试运行，正推进真原生化以满足 App Store 商业化上架要求。", 14, INK_DIM, False)]],
         line_spacing=1.3)

headers = ["层级", "技术方案"]
rows = [
    ["后端", "Go 语言编写，无外部数据库依赖，本地 JSON 原子写入持久化，服务器成本极低"],
    ["数据源", "同时对接 Semantic Scholar、arXiv、OpenAlex 三大权威学术源，自动降级容错"],
    ["AI 能力", "可插拔 OpenAI-compatible LLM 接口（当前对接智谱 GLM），无 Key 时自动降级为本地证据抽取，保证功能永远可用"],
    ["前端", "零构建原生 H5，移动端手势深度打磨（滑动判断、长按拖拽、防误触/防选中）"],
    ["客户端化", "Capacitor 混合壳 + 原生 Haptics 震动反馈，具备向 iOS/Android 原生 App 演进的完整路径"],
]
table_slide(s, headers, rows, Inches(0.75), Inches(2.95), Inches(11.83), Inches(3.6),
            col_widths=[2, 8])

# =====================================================================
# 11. 差异化定位
# =====================================================================
s = new_slide()
add_kicker(s, "竞争格局")
add_title(s, "唯一把「发现→判断→归档→阅读规划→热点追踪」做成一体化闭环的产品")

headers = ["类型", "代表产品", "局限"]
rows = [
    ["学术搜索引擎", "Google Scholar、Semantic Scholar 官网", "仍是关键词搜索范式，判断效率完全依赖人工"],
    ["文献管理软件", "Zotero、Mendeley、EndNote", "擅长管理已收藏文献，不解决「发现」与「何时读」问题"],
    ["滑动式发现工具", "Zinote（个人开发者独立产品）", "仅实现「按话题抓取摘要+滑动」单点功能，无 Library / Schedule / Trending 等配套生态，停留在 Demo 层面"],
    ["PaperSwipe", "—", "发现 + 判断 + 归档 + 阅读规划 + 热点追踪，移动端一体化闭环"],
]
table_slide(s, headers, rows, Inches(0.75), Inches(1.95), Inches(11.83), Inches(3.6),
            col_widths=[2.2, 3, 6.6])

add_text(s, Inches(0.75), Inches(5.85), Inches(11.3), Inches(1.0),
         [[("现有产品之间高度割裂——用户需要同时打开 Google Scholar 搜索、翻译工具理解摘要、", 13, INK_DIM, False)],
          [("Zotero 管理收藏、Twitter 追踪热点，PaperSwipe 用一个 App 把这条链路完全打通。", 13, INK_DIM, False)]],
         line_spacing=1.3)

# =====================================================================
# 12. 市场机会
# =====================================================================
s = new_slide()
add_kicker(s, "市场机会")
add_title(s, "一个规模足够大、且持续增长的市场")

stats = [
    ("880万+", "全球科研人员数量\n（UNESCO 数据）", VIOLET),
    ("388万+", "中国在读研究生规模\n（教育部 2023 年数据）", CYAN),
    ("20万+/年", "arXiv 年新增论文提交量\n较十年前翻倍以上", CORAL),
]
card_w = Inches(3.6)
gap = Inches(0.35)
left0 = Inches(0.9)
top = Inches(2.1)
for i, (num, label, color) in enumerate(stats):
    x = left0 + (card_w + gap) * i
    add_rect(s, x, top, card_w, Inches(2.3), WHITE, radius=0.08, line=True, line_color=LINE)
    add_text(s, x, top + Inches(0.4), card_w, Inches(0.8),
             [[(num, 32, color, True)]], align=PP_ALIGN.CENTER)
    lines = label.split("\n")
    add_text(s, x + Inches(0.3), top + Inches(1.35), card_w - Inches(0.6), Inches(0.8),
             [[(ln, 12.5, INK_DIM, False)] for ln in lines], align=PP_ALIGN.CENTER, line_spacing=1.25)

add_text(s, Inches(0.9), Inches(4.85), Inches(11.2), Inches(1.6),
         [[('"高频读论文、却缺乏高效发现工具"是一个规模足够大、且仍在持续增长的市场。', 15, INK, True)],
          [("PaperSwipe 计划优先切入中文科研用户群体，验证付费与留存后再逐步向英文市场扩展。", 14, INK_DIM, False)]],
         line_spacing=1.4)

# =====================================================================
# 13. 早期验证
# =====================================================================
s = new_slide()
add_kicker(s, "早期验证")
add_title(s, "真实用户已经在告诉我们：这是一个真痛点")

bullets(s, [
    "黑客松路演现场收到大量用户正向认可：反复被验证的核心反馈是「这是一个真实痛点、我确实需要这样一个工具」",
    "用户普遍认可「让枯燥的科研过程娱乐化、年轻化」这一产品理念",
    "正在运营小红书账号，持续发布产品内容与使用场景，已收到大量积极反馈",
    "验证了「科研 + 娱乐化交互」这一心智在真实用户中具备传播力，为规模化获客积累了初步的内容与渠道基础",
], Inches(0.75), Inches(2.0), Inches(6.9), Inches(4.5), size=14, gap=16)

picture_fit(s, os.path.join(ASSETS, "小红书.jpg"), Inches(8.15), Inches(1.5), Inches(2.4), Inches(5.2))
picture_fit(s, os.path.join(ASSETS, "微信群.jpg"), Inches(10.6), Inches(1.5), Inches(2.4), Inches(5.2))

# =====================================================================
# 14. 商业模式
# =====================================================================
s = new_slide()
add_kicker(s, "商业模式")
add_title(s, "Freemium 订阅制为主线，多层收入曲线")

headers = ["收入来源", "目标用户", "说明"]
rows = [
    ["Pro 订阅（核心）", "高频研究生 / 科研人员", "无限量发现、AI 阅读规划与对话助手、更高质量摘要，App Store/Google Play 标准订阅制内购"],
    ["机构/实验室版", "高校实验室、企业研发团队", "团队共享文献库、管理者仪表盘，按坐席数订阅"],
    ["数据与 API 授权", "出版方、文献管理工具", "基于热点追踪能力，输出趋势报告或开放检索/推荐 API"],
    ["生态联动（远期）", "Zotero 等工具、学术会议/期刊", "导出/联动场景中的轻量级合作分佣或联合推广"],
]
table_slide(s, headers, rows, Inches(0.75), Inches(1.95), Inches(11.83), Inches(3.3),
            col_widths=[2.2, 3, 6.6])

add_text(s, Inches(0.75), Inches(5.5), Inches(11.3), Inches(1.4),
         [[("免费层保证核心体验可用，作为获客与口碑传播基本盘；Pro 订阅层解锁无限量发现、", 13, INK_DIM, False)],
          [("智能排期与对话助手、多设备同步等——高频刚需 + 按量计费成本结构，是健康的商业模式。", 13, INK_DIM, False)]],
         line_spacing=1.3)

# =====================================================================
# 15. 目前进展 & Roadmap
# =====================================================================
s = new_slide()
add_kicker(s, "目前进展 & 未来计划")
add_title(s, "MVP 已跑通，下一步是规模化增长")

add_text(s, Inches(0.75), Inches(1.85), Inches(2.2), Inches(0.4), [[("目前进展", 15, INK, True)]])
bullets(s, [
    "核心 MVP 全功能闭环已上线（四大模块全部跑通）",
    "支持标签筛选、BibTeX/Zotero 导出",
    "移动端交互细节大量打磨，产品完成度达可对外测试成熟度",
    "完成 iOS 端 Capacitor 封装，可在真机调试运行",
    "已启动小红书/微信社群早期运营",
], Inches(0.75), Inches(2.3), Inches(5.8), Inches(4.4), size=12.5, gap=10, mark="✓", mark_color=GREEN)

headers = ["阶段", "关键目标"]
rows = [
    ["Phase 1·已完成", "核心 MVP 全功能闭环，早期种子用户验证"],
    ["Phase 2·0~6个月", "账号体系与多租户后端；iOS 真原生化上架；接入订阅制内购"],
    ["Phase 3·6~12个月", "Android 原生 App；PDF 全文解析；机构/实验室团队版；开放 API"],
    ["Phase 4·12个月+", "海外市场扩张；出版方数据合作；领域热点报告等数据增值服务"],
]
table_slide(s, headers, rows, Inches(6.85), Inches(1.85), Inches(5.75), Inches(4.9),
            col_widths=[2, 4.8], body_size=11.5, header_size=12.5)

# =====================================================================
# 16. 团队介绍
# =====================================================================
s = new_slide()
add_kicker(s, "团队介绍")
add_title(s, "小而精，全栈能力覆盖完整")

roles = [
    ("产品与前端", "产品定义、功能与交互设计、H5 前端开发、iOS 原生化落地", VIOLET),
    ("后端与数据库开发", "Go 后端架构、多学术源检索与自动降级、AI 摘要与推荐策略、数据存储与状态管理", CYAN),
    ("运营宣传", "用户增长、社区运营（小红书/微信社群）、产品宣传物料与路演、用户反馈收集与闭环", CORAL),
]
card_w = Inches(3.75)
gap = Inches(0.3)
left0 = Inches(0.75)
top = Inches(2.15)
card_h = Inches(3.4)
for i, (name, desc, color) in enumerate(roles):
    x = left0 + (card_w + gap) * i
    add_rect(s, x, top, card_w, card_h, WHITE, radius=0.06, line=True, line_color=LINE)
    add_rect(s, x + Inches(0.3), top + Inches(0.35), Inches(0.55), Inches(0.55), color, radius=1.0)
    add_text(s, x + Inches(0.3), top + Inches(1.15), card_w - Inches(0.6), Inches(0.6),
             [[(name, 16.5, INK, True)]])
    add_text(s, x + Inches(0.3), top + Inches(1.75), card_w - Inches(0.6), Inches(1.5),
             [[(desc, 12, INK_DIM, False)]], line_spacing=1.3)

add_text(s, Inches(0.75), Inches(5.85), Inches(11.3), Inches(1.0),
         [[("团队成员均有扎实工程背景与真实科研/读论文经历——PaperSwipe 本质上是团队", 13, INK_DIM, False)],
          [("「为自己造工具」后发现的强共性需求，这也是我们对 PMF 有信心的原因之一。", 13, INK_DIM, False)]],
         line_spacing=1.3)

# =====================================================================
# 17. 融资需求
# =====================================================================
s = new_slide()
add_kicker(s, "融资需求")
add_title(s, "启动天使轮/种子轮，从 MVP 迈向规模化增长")

headers = ["用途", "占比", "说明"]
rows = [
    ["研发投入", "40%", "多租户后端、iOS/Android 原生化、订阅内购接入、PDF 全文解析"],
    ["服务器与 API 成本", "20%", "LLM 摘要调用、学术源检索、图片生成等按量计费成本"],
    ["市场推广与获客", "25%", "高校社群运营、内容营销、KOL 合作、应用商店 ASO"],
    ["团队扩编", "10%", "优先补强后端工程与增长/运营岗位"],
    ["合规与其他", "5%", "应用商店合规审核、法务、财务等基础运营开支"],
]
table_slide(s, headers, rows, Inches(0.75), Inches(1.95), Inches(11.83), Inches(3.9),
            col_widths=[3, 1.4, 6.5])

add_text(s, Inches(0.75), Inches(6.15), Inches(11.3), Inches(1.0),
         [[("里程碑：iOS App 正式上架 App Store；Pro 订阅转化率达到健康水平；", 13, INK_DIM, False)],
          [("建立稳定次周/次月留存曲线；完成至少 1 个高校实验室/企业机构版试点合作。", 13, INK_DIM, False)]],
         line_spacing=1.3)

# =====================================================================
# 18. 结尾 / 联系方式
# =====================================================================
s = new_slide(bg=DARK_PANEL)
add_text(s, Inches(1.0), Inches(0.9), Inches(11), Inches(1.0),
         [[("把「找论文」的力气，还给「读论文」本身。", 26, WHITE, True)]])
add_text(s, Inches(1.0), Inches(1.55), Inches(10.5), Inches(0.8),
         [[("传统方式：搜索、翻页、翻译、收藏、然后忘记。", 14, RGBColor(0xC7,0xC9,0xD3), False)],
          [("PaperSwipe：AI 帮你发现、5 秒帮你判断、自动帮你安排。", 14, WHITE, True)]],
         line_spacing=1.35)

qrs = [
    (os.path.join(ASSETS, "小红书.jpg"), "小红书", "关注产品更新与幕后花絮"),
    (os.path.join(ASSETS, "微信群.jpg"), "用户群", "加入用户群一起聊 idea"),
    (os.path.join(ASSETS, "微信.jpg"), "Creator 微信", "合作/内测申请，24 小时内回复"),
]
card_w = Inches(3.5)
gap = Inches(0.4)
left0 = Inches(1.0)
top = Inches(2.7)
for i, (path, name, desc) in enumerate(qrs):
    x = left0 + (card_w + gap) * i
    add_rect(s, x, top, card_w, Inches(4.2), RGBColor(0x22,0x23,0x30), radius=0.06)
    picture_fit(s, path, x + Inches(0.55), top + Inches(0.35), Inches(2.4), Inches(2.4), border=False)
    add_text(s, x, top + Inches(2.95), card_w, Inches(0.5),
             [[(name, 16, WHITE, True)]], align=PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.3), top + Inches(3.45), card_w - Inches(0.6), Inches(0.7),
             [[(desc, 11.5, RGBColor(0xA0,0xA3,0xB0), False)]], align=PP_ALIGN.CENTER, line_spacing=1.3)

add_text(s, Inches(1.0), SH - Inches(0.7), Inches(8), Inches(0.4),
         [[("© PaperSwipe — Discover, not Search.", 11, RGBColor(0x8A,0x8E,0x9C), False)]])

prs.save(OUT)
print("Saved:", OUT)
